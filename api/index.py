import csv
import json
import mimetypes
import re
from datetime import datetime
from http.server import BaseHTTPRequestHandler
from io import BytesIO
from pathlib import Path
from urllib.parse import parse_qs, unquote, urlparse, urljoin
from urllib.request import Request, urlopen

from openpyxl import load_workbook

from ctd_model import calculate

CME_PAGE = "https://www.cmegroup.com/trading/interest-rates/treasury-conversion-factors.html"
ROOT = Path(__file__).resolve().parents[1]


def _date(value):
    return value.strftime("%Y-%m-%d") if isinstance(value, datetime) else str(value)[:10]


def _load_basket(month):
    try:
        page_request = Request(CME_PAGE, headers={"User-Agent": "Mozilla/5.0 CTD-calculator"})
        with urlopen(page_request, timeout=20) as response:
            page = response.read().decode("utf-8", errors="replace")
        match = re.search(r"href=[\"']([^\"']*TCF\.xlsx[^\"']*)[\"']", page, flags=re.I)
        if not match:
            raise RuntimeError("CME workbook link not found")
        workbook_url = urljoin(CME_PAGE, match.group(1).replace("&amp;", "&"))
        request = Request(workbook_url, headers={"User-Agent": "Mozilla/5.0 CTD-calculator"})
        with urlopen(request, timeout=45) as response:
            workbook_bytes = response.read()
        sheet = load_workbook(BytesIO(workbook_bytes), data_only=True, read_only=True)["Conversion Factors"]
        target = datetime.strptime(month + "-01", "%Y-%m-%d")
        section = next(row[0].row for row in sheet.iter_rows(min_col=1, max_col=1)
                       if str(row[0].value).strip() == "U.S. TREASURY BOND FUTURES CONTRACT")
        target_col = None
        header_row = None
        for row_num in range(section + 1, min(section + 10, sheet.max_row) + 1):
            for col_num in range(1, sheet.max_column + 1):
                value = sheet.cell(row_num, col_num).value
                if isinstance(value, datetime) and value.year == target.year and value.month == target.month:
                    header_row, target_col = row_num, col_num
                    break
            if target_col:
                break
        if target_col is None:
            raise ValueError("CME workbook has no ZB factors for " + month)
        records = []
        for row_num in range(header_row + 1, sheet.max_row + 1):
            first = sheet.cell(row_num, 1).value
            if isinstance(first, str) and "FUTURES CONTRACT" in first:
                break
            cusip = sheet.cell(row_num, 6).value
            factor = sheet.cell(row_num, target_col).value
            if isinstance(cusip, str) and cusip.strip().isdigit() and len(cusip.strip()) == 9 and isinstance(factor, (int, float)):
                records.append({"cusip":cusip.strip(), "coupon_pct":float(sheet.cell(row_num, 3).value),
                                "issue_date":_date(sheet.cell(row_num, 4).value),
                                "maturity":_date(sheet.cell(row_num, 5).value),
                                "conversion_factor":float(factor)})
        if not records:
            raise ValueError("CME workbook returned no eligible ZB bonds")
        return {"basket":records, "source":"CME Group workbook", "month":month}
    except Exception as exc:
        fallback = ROOT / "Treasury_Bond_Futures_ZB_Dec2026_Basket_2026-10-08.csv"
        if month == "2026-12" and fallback.exists():
            with fallback.open(newline="") as stream:
                records = list(csv.DictReader(stream))
            for bond in records:
                bond["coupon_pct"] = float(bond["coupon_pct"])
                bond["conversion_factor"] = float(bond["conversion_factor"])
            return {"basket":records, "source":"CME cached 2026-10-08 (refresh error: " + str(exc) + ")", "month":month}
        raise


def _load_market():
    url = "https://fred.stlouisfed.org/graph/fredgraph.csv?id=DGS10,DGS20,DGS30&cosd=2000-01-01"
    with urlopen(url, timeout=30) as response:
        text = response.read().decode("utf-8")
    points = []
    for row in csv.DictReader(text.splitlines()):
        try:
            point = {"date":row.get("observation_date", row.get("DATE")),
                     "dgs10":float(row["DGS10"]), "dgs20":float(row["DGS20"]), "dgs30":float(row["DGS30"])}
            if point["date"]:
                points.append(point)
        except (TypeError, ValueError, KeyError):
            continue
    if not points:
        raise RuntimeError("FRED returned no complete 10/20/30-year yield observations")
    return {"source":"FRED / U.S. Treasury H.15", "as_of":points[-1]["date"], "points":points}


class handler(BaseHTTPRequestHandler):
    STATIC_FILES = {
        "/": "index.html",
        "/index.html": "index.html",
        "/styles.css": "styles.css",
        "/app.js": "app.js",
        "/Convexity-Maven-Fat-Cat.pdf": "Convexity-Maven-Fat-Cat.pdf",
        "/Treasury_Bond_Futures_CTD_Delivery_Option_Calculator.ipynb": "Treasury_Bond_Futures_CTD_Delivery_Option_Calculator.ipynb",
        "/CUSIP_Yield_Quotes_Template.csv": "CUSIP_Yield_Quotes_Template.csv",
        "/Treasury_Bond_Futures_ZB_Dec2026_Basket_2026-10-08.csv": "Treasury_Bond_Futures_ZB_Dec2026_Basket_2026-10-08.csv",
    }

    def _send(self, status, value, cache="no-store"):
        body = json.dumps(value, allow_nan=False).encode("utf-8")
        self.send_response(status)
        self.send_header("Content-Type", "application/json; charset=utf-8")
        self.send_header("Cache-Control", cache)
        self.send_header("Access-Control-Allow-Origin", "*")
        self.send_header("Access-Control-Allow-Methods", "GET, POST, OPTIONS")
        self.send_header("Access-Control-Allow-Headers", "Content-Type")
        self.end_headers()
        self.wfile.write(body)

    def _serve_static(self):
        path = unquote(urlparse(self.path).path)
        filename = self.STATIC_FILES.get(path)
        if filename is None:
            return self._send(404, {"error": "Not found"})
        file_path = ROOT / filename
        try:
            body = file_path.read_bytes()
        except FileNotFoundError:
            return self._send(404, {"error": "Not found"})
        self.send_response(200)
        self.send_header("Content-Type", mimetypes.guess_type(filename)[0] or "application/octet-stream")
        self.send_header("Content-Length", str(len(body)))
        self.send_header("Cache-Control", "public, max-age=3600")
        self.end_headers()
        self.wfile.write(body)

    def do_HEAD(self):
        path = unquote(urlparse(self.path).path)
        filename = self.STATIC_FILES.get(path)
        if filename is None:
            return self._send(404, {"error": "Not found"})
        file_path = ROOT / filename
        try:
            size = file_path.stat().st_size
        except FileNotFoundError:
            return self._send(404, {"error": "Not found"})
        self.send_response(200)
        self.send_header("Content-Type", mimetypes.guess_type(filename)[0] or "application/octet-stream")
        self.send_header("Content-Length", str(size))
        self.send_header("Cache-Control", "public, max-age=3600")
        self.end_headers()

    def do_OPTIONS(self):
        self._send(204, {})

    def do_GET(self):
        query = parse_qs(urlparse(self.path).query)
        route = query.get("route", [""])[0]
        try:
            if route == "basket":
                month = query.get("month", ["2026-12"])[0]
                return self._send(200, _load_basket(month), "s-maxage=3600, stale-while-revalidate=86400")
            if route == "market":
                return self._send(200, _load_market(), "s-maxage=1800, stale-while-revalidate=3600")
            if not urlparse(self.path).path.startswith("/api/"):
                return self._serve_static()
            return self._send(404, {"error":"Unknown API route"})
        except Exception as exc:
            return self._send(502, {"error":str(exc)})

    def do_POST(self):
        query = parse_qs(urlparse(self.path).query)
        if query.get("route", [""])[0] != "calculate":
            return self._send(404, {"error":"Unknown API route"})
        try:
            length = int(self.headers.get("Content-Length", "0"))
            payload = json.loads(self.rfile.read(length))
            return self._send(200, calculate(payload))
        except Exception as exc:
            return self._send(400, {"error":str(exc)})
