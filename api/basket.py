import csv
import json
import re
from datetime import datetime
from http.server import BaseHTTPRequestHandler
from io import BytesIO
from pathlib import Path
from urllib.parse import parse_qs, urlparse, urljoin
from urllib.request import Request, urlopen

from openpyxl import load_workbook

CME_PAGE = "https://www.cmegroup.com/trading/interest-rates/treasury-conversion-factors.html"
ROOT = Path(__file__).resolve().parents[1]


def _date(value):
    return value.strftime("%Y-%m-%d") if isinstance(value, datetime) else str(value)[:10]


def load_basket(month):
    try:
        page_req = Request(CME_PAGE, headers={"User-Agent": "Mozilla/5.0 CTD-calculator"})
        with urlopen(page_req, timeout=20) as response:
            page = response.read().decode("utf-8", errors="replace")
        match = re.search(r"href=[\"']([^\"']*TCF\.xlsx[^\"']*)[\"']", page, flags=re.I)
        if not match:
            raise RuntimeError("CME workbook link not found")
        workbook_url = urljoin(CME_PAGE, match.group(1).replace("&amp;", "&"))
        req = Request(workbook_url, headers={"User-Agent": "Mozilla/5.0 CTD-calculator"})
        with urlopen(req, timeout=45) as response:
            workbook_bytes = response.read()
        sheet = load_workbook(BytesIO(workbook_bytes), data_only=True, read_only=True)["Conversion Factors"]
        target = datetime.strptime(month + "-01", "%Y-%m-%d")
        section = next(row[0].row for row in sheet.iter_rows(min_col=1, max_col=1)
                       if str(row[0].value).strip() == "U.S. TREASURY BOND FUTURES CONTRACT")
        target_col = None
        header_row = None
        for row_num in range(section + 1, min(section + 10, sheet.max_row) + 1):
            for col_num in range(1, sheet.max_column + 1):
                value = sheet.cell(row_num, col_num).value
                if isinstance(value, datetime) and value.year == target.year and value.month == target.month:
                    header_row, target_col = row_num, col_num
                    break
            if target_col:
                break
        if target_col is None:
            raise ValueError("CME workbook has no ZB factors for " + month)
        records = []
        for row_num in range(header_row + 1, sheet.max_row + 1):
            first = sheet.cell(row_num, 1).value
            if isinstance(first, str) and "FUTURES CONTRACT" in first:
                break
            cusip = sheet.cell(row_num, 6).value
            cf = sheet.cell(row_num, target_col).value
            if isinstance(cusip, str) and cusip.strip().isdigit() and len(cusip.strip()) == 9 and isinstance(cf, (int, float)):
                records.append({
                    "cusip": cusip.strip(),
                    "coupon_pct": float(sheet.cell(row_num, 3).value),
                    "issue_date": _date(sheet.cell(row_num, 4).value),
                    "maturity": _date(sheet.cell(row_num, 5).value),
                    "conversion_factor": float(cf),
                })
        if not records:
            raise ValueError("CME workbook returned no eligible ZB bonds")
        return {"basket": records, "source": "CME Group workbook", "month": month}
    except Exception as exc:
        fallback = ROOT / "Treasury_Bond_Futures_ZB_Dec2026_Basket_2026-10-08.csv"
        if month == "2026-12" and fallback.exists():
            with fallback.open(newline="") as stream:
                records = list(csv.DictReader(stream))
            for row in records:
                row["coupon_pct"] = float(row["coupon_pct"])
                row["conversion_factor"] = float(row["conversion_factor"])
            return {"basket": records, "source": "CME cached 2026-10-08 (refresh error: " + str(exc) + ")", "month": month}
        raise


class handler(BaseHTTPRequestHandler):
    def do_GET(self):
        params = parse_qs(urlparse(self.path).query)
        month = params.get("month", ["2026-12"])[0]
        try:
            result = load_basket(month)
            body = json.dumps(result).encode("utf-8")
            self.send_response(200)
            self.send_header("Content-Type", "application/json; charset=utf-8")
            self.send_header("Cache-Control", "s-maxage=3600, stale-while-revalidate=86400")
            self.send_header("Access-Control-Allow-Origin", "*")
            self.end_headers()
            self.wfile.write(body)
        except Exception as exc:
            self.send_response(400)
            self.send_header("Content-Type", "application/json; charset=utf-8")
            self.send_header("Access-Control-Allow-Origin", "*")
            self.end_headers()
            self.wfile.write(json.dumps({"error": str(exc)}).encode("utf-8"))
